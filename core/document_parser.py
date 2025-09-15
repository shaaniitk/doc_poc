"""Enhanced Document Parser for LangGraph Integration

This module provides a modern, async-capable document parser that integrates
seamlessly with the LangGraph orchestrator and centralized state management.
It supports multiple document formats with intelligent format detection and
robust error handling.
"""

import asyncio
import logging
from typing import Dict, List, Optional, Any, Union, BinaryIO, TextIO
from pathlib import Path
from dataclasses import dataclass, field
from enum import Enum, auto
import mimetypes
import chardet
from datetime import datetime, timezone

# Document processing imports
try:
    import PyPDF2
    import pdfplumber
    HAS_PDF_SUPPORT = True
except ImportError:
    HAS_PDF_SUPPORT = False
    logging.warning("PDF support not available. Install PyPDF2 and pdfplumber for PDF processing.")

try:
    from docx import Document as DocxDocument
    HAS_DOCX_SUPPORT = True
except ImportError:
    HAS_DOCX_SUPPORT = False
    logging.warning("DOCX support not available. Install python-docx for Word document processing.")

try:
    import markdown
    HAS_MARKDOWN_SUPPORT = True
except ImportError:
    HAS_MARKDOWN_SUPPORT = False
    logging.warning("Markdown support not available. Install markdown for .md file processing.")

logger = logging.getLogger(__name__)


class DocumentFormat(Enum):
    """Supported document formats."""
    PDF = auto()
    DOCX = auto()
    TXT = auto()
    MARKDOWN = auto()
    TEX = auto()
    HTML = auto()
    RTF = auto()
    UNKNOWN = auto()


class ParsingStrategy(Enum):
    """Different parsing strategies for various document types."""
    FAST = auto()          # Quick parsing, may miss some formatting
    BALANCED = auto()      # Balance between speed and accuracy
    COMPREHENSIVE = auto() # Thorough parsing, slower but more accurate
    CUSTOM = auto()        # Custom parsing logic


@dataclass
class DocumentMetadata:
    """Comprehensive document metadata."""
    file_path: Path
    file_size: int
    format: DocumentFormat
    encoding: Optional[str] = None
    mime_type: Optional[str] = None
    creation_time: Optional[datetime] = None
    modification_time: Optional[datetime] = None
    page_count: Optional[int] = None
    word_count: Optional[int] = None
    character_count: Optional[int] = None
    language: Optional[str] = None
    title: Optional[str] = None
    author: Optional[str] = None
    subject: Optional[str] = None
    parsing_strategy: ParsingStrategy = ParsingStrategy.BALANCED
    parsing_time: Optional[float] = None
    quality_score: Optional[float] = None
    warnings: List[str] = field(default_factory=list)
    custom_properties: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    """Container for parsed document content and metadata."""
    content: str
    metadata: DocumentMetadata
    sections: List[Dict[str, Any]] = field(default_factory=list)
    tables: List[Dict[str, Any]] = field(default_factory=list)
    images: List[Dict[str, Any]] = field(default_factory=list)
    links: List[Dict[str, Any]] = field(default_factory=list)
    footnotes: List[Dict[str, Any]] = field(default_factory=list)
    headers_footers: Dict[str, str] = field(default_factory=dict)
    raw_data: Optional[Any] = None


class DocumentParser:
    """Enhanced document parser with multi-format support and async capabilities."""
    
    def __init__(self, 
                 default_strategy: ParsingStrategy = ParsingStrategy.BALANCED,
                 enable_metadata_extraction: bool = True,
                 enable_structure_analysis: bool = True,
                 max_file_size_mb: int = 100):
        self.default_strategy = default_strategy
        self.enable_metadata_extraction = enable_metadata_extraction
        self.enable_structure_analysis = enable_structure_analysis
        self.max_file_size_mb = max_file_size_mb
        
        # Format detection mappings
        self.format_mappings = {
            '.pdf': DocumentFormat.PDF,
            '.docx': DocumentFormat.DOCX,
            '.doc': DocumentFormat.DOCX,
            '.txt': DocumentFormat.TXT,
            '.md': DocumentFormat.MARKDOWN,
            '.markdown': DocumentFormat.MARKDOWN,
            '.tex': DocumentFormat.TEX,
            '.latex': DocumentFormat.TEX,
            '.html': DocumentFormat.HTML,
            '.htm': DocumentFormat.HTML,
            '.rtf': DocumentFormat.RTF
        }
        
        logger.info(f"DocumentParser initialized with strategy: {default_strategy.name}")
    
    async def parse_document_async(self, file_path: Union[str, Path], 
                                 strategy: Optional[ParsingStrategy] = None) -> ParsedDocument:
        """Asynchronously parse a document."""
        return await asyncio.get_event_loop().run_in_executor(
            None, self.parse_document, file_path, strategy
        )
    
    def parse_document(self, file_path: Union[str, Path], 
                      strategy: Optional[ParsingStrategy] = None) -> ParsedDocument:
        """Parse a document with comprehensive error handling and format detection."""
        start_time = datetime.now(timezone.utc)
        file_path = Path(file_path)
        strategy = strategy or self.default_strategy
        
        try:
            # Validate file
            self._validate_file(file_path)
            
            # Detect format
            document_format = self._detect_format(file_path)
            
            # Create metadata
            metadata = self._create_metadata(file_path, document_format, strategy)
            
            # Parse based on format
            parsed_doc = self._parse_by_format(file_path, document_format, strategy, metadata)
            
            # Post-processing
            if self.enable_structure_analysis:
                self._analyze_structure(parsed_doc)
            
            # Calculate quality score
            parsed_doc.metadata.quality_score = self._calculate_quality_score(parsed_doc)
            
            # Record parsing time
            parsing_time = (datetime.now(timezone.utc) - start_time).total_seconds()
            parsed_doc.metadata.parsing_time = parsing_time
            
            logger.info(f"Successfully parsed {file_path.name} ({document_format.name}) in {parsing_time:.2f}s")
            return parsed_doc
            
        except Exception as e:
            logger.error(f"Failed to parse document {file_path}: {e}")
            raise DocumentParsingError(f"Parsing failed for {file_path}: {e}") from e
    
    def _validate_file(self, file_path: Path) -> None:
        """Validate file exists and is within size limits."""
        if not file_path.exists():
            raise FileNotFoundError(f"Document not found: {file_path}")
        
        if not file_path.is_file():
            raise ValueError(f"Path is not a file: {file_path}")
        
        file_size_mb = file_path.stat().st_size / (1024 * 1024)
        if file_size_mb > self.max_file_size_mb:
            raise ValueError(f"File too large: {file_size_mb:.1f}MB > {self.max_file_size_mb}MB")
    
    def _detect_format(self, file_path: Path) -> DocumentFormat:
        """Detect document format using extension and MIME type."""
        # Try extension first
        extension = file_path.suffix.lower()
        if extension in self.format_mappings:
            detected_format = self.format_mappings[extension]
            logger.debug(f"Format detected by extension: {detected_format.name}")
            return detected_format
        
        # Try MIME type detection
        mime_type, _ = mimetypes.guess_type(str(file_path))
        if mime_type:
            mime_mappings = {
                'application/pdf': DocumentFormat.PDF,
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document': DocumentFormat.DOCX,
                'application/msword': DocumentFormat.DOCX,
                'text/plain': DocumentFormat.TXT,
                'text/markdown': DocumentFormat.MARKDOWN,
                'text/x-tex': DocumentFormat.TEX,
                'text/html': DocumentFormat.HTML,
                'application/rtf': DocumentFormat.RTF
            }
            
            if mime_type in mime_mappings:
                detected_format = mime_mappings[mime_type]
                logger.debug(f"Format detected by MIME type: {detected_format.name}")
                return detected_format
        
        logger.warning(f"Unknown format for {file_path}, defaulting to TXT")
        return DocumentFormat.UNKNOWN
    
    def _create_metadata(self, file_path: Path, document_format: DocumentFormat, 
                        strategy: ParsingStrategy) -> DocumentMetadata:
        """Create comprehensive document metadata."""
        stat = file_path.stat()
        mime_type, _ = mimetypes.guess_type(str(file_path))
        
        return DocumentMetadata(
            file_path=file_path,
            file_size=stat.st_size,
            format=document_format,
            mime_type=mime_type,
            creation_time=datetime.fromtimestamp(stat.st_ctime, timezone.utc),
            modification_time=datetime.fromtimestamp(stat.st_mtime, timezone.utc),
            parsing_strategy=strategy
        )
    
    def _parse_by_format(self, file_path: Path, document_format: DocumentFormat, 
                        strategy: ParsingStrategy, metadata: DocumentMetadata) -> ParsedDocument:
        """Parse document based on detected format."""
        parsers = {
            DocumentFormat.PDF: self._parse_pdf,
            DocumentFormat.DOCX: self._parse_docx,
            DocumentFormat.TXT: self._parse_text,
            DocumentFormat.MARKDOWN: self._parse_markdown,
            DocumentFormat.TEX: self._parse_tex,
            DocumentFormat.HTML: self._parse_html,
            DocumentFormat.RTF: self._parse_rtf,
            DocumentFormat.UNKNOWN: self._parse_text  # Fallback to text
        }
        
        parser_func = parsers.get(document_format, self._parse_text)
        return parser_func(file_path, strategy, metadata)
    
    def _parse_pdf(self, file_path: Path, strategy: ParsingStrategy, 
                  metadata: DocumentMetadata) -> ParsedDocument:
        """Parse PDF documents with multiple extraction methods."""
        if not HAS_PDF_SUPPORT:
            raise DocumentParsingError("PDF support not available. Install PyPDF2 and pdfplumber.")
        
        content = ""
        sections = []
        tables = []
        
        try:
            # Use pdfplumber for better text extraction
            import pdfplumber
            
            with pdfplumber.open(file_path) as pdf:
                metadata.page_count = len(pdf.pages)
                
                for page_num, page in enumerate(pdf.pages, 1):
                    page_text = page.extract_text() or ""
                    content += page_text + "\n\n"
                    
                    # Extract tables if comprehensive strategy
                    if strategy == ParsingStrategy.COMPREHENSIVE:
                        page_tables = page.extract_tables()
                        for table in page_tables or []:
                            tables.append({
                                "page": page_num,
                                "data": table,
                                "type": "table"
                            })
                    
                    sections.append({
                        "type": "page",
                        "number": page_num,
                        "content": page_text,
                        "word_count": len(page_text.split()) if page_text else 0
                    })
            
        except Exception as e:
            # Fallback to PyPDF2
            logger.warning(f"pdfplumber failed, trying PyPDF2: {e}")
            
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                metadata.page_count = len(pdf_reader.pages)
                
                for page_num, page in enumerate(pdf_reader.pages, 1):
                    page_text = page.extract_text()
                    content += page_text + "\n\n"
                    
                    sections.append({
                        "type": "page",
                        "number": page_num,
                        "content": page_text,
                        "word_count": len(page_text.split()) if page_text else 0
                    })
        
        return ParsedDocument(
            content=content.strip(),
            metadata=metadata,
            sections=sections,
            tables=tables
        )
    
    def _parse_docx(self, file_path: Path, strategy: ParsingStrategy, 
                   metadata: DocumentMetadata) -> ParsedDocument:
        """Parse DOCX documents with structure preservation."""
        if not HAS_DOCX_SUPPORT:
            raise DocumentParsingError("DOCX support not available. Install python-docx.")
        
        from docx import Document as DocxDocument
        
        doc = DocxDocument(file_path)
        content = ""
        sections = []
        tables = []
        
        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                content += para.text + "\n"
                
                sections.append({
                    "type": "paragraph",
                    "content": para.text,
                    "style": para.style.name if para.style else None
                })
        
        # Extract tables if enabled
        if strategy in [ParsingStrategy.BALANCED, ParsingStrategy.COMPREHENSIVE]:
            for table_idx, table in enumerate(doc.tables):
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                
                tables.append({
                    "index": table_idx,
                    "data": table_data,
                    "type": "table"
                })
        
        # Extract document properties
        if self.enable_metadata_extraction and doc.core_properties:
            props = doc.core_properties
            metadata.title = props.title
            metadata.author = props.author
            metadata.subject = props.subject
        
        return ParsedDocument(
            content=content.strip(),
            metadata=metadata,
            sections=sections,
            tables=tables
        )
    
    def _parse_text(self, file_path: Path, strategy: ParsingStrategy, 
                   metadata: DocumentMetadata) -> ParsedDocument:
        """Parse plain text files with encoding detection."""
        # Detect encoding
        with open(file_path, 'rb') as file:
            raw_data = file.read()
            encoding_result = chardet.detect(raw_data)
            encoding = encoding_result.get('encoding', 'utf-8')
            metadata.encoding = encoding
        
        # Read with detected encoding
        try:
            with open(file_path, 'r', encoding=encoding) as file:
                content = file.read()
        except UnicodeDecodeError:
            # Fallback to utf-8 with error handling
            with open(file_path, 'r', encoding='utf-8', errors='replace') as file:
                content = file.read()
                metadata.warnings.append("Encoding issues detected, some characters may be corrupted")
        
        # Simple section detection for text files
        sections = []
        if strategy in [ParsingStrategy.BALANCED, ParsingStrategy.COMPREHENSIVE]:
            lines = content.split('\n')
            current_section = []
            
            for line in lines:
                if line.strip() == "" and current_section:
                    # End of section
                    section_content = '\n'.join(current_section)
                    sections.append({
                        "type": "section",
                        "content": section_content,
                        "line_count": len(current_section)
                    })
                    current_section = []
                elif line.strip():
                    current_section.append(line)
            
            # Add final section
            if current_section:
                section_content = '\n'.join(current_section)
                sections.append({
                    "type": "section",
                    "content": section_content,
                    "line_count": len(current_section)
                })
        
        return ParsedDocument(
            content=content,
            metadata=metadata,
            sections=sections
        )
    
    def _parse_markdown(self, file_path: Path, strategy: ParsingStrategy, 
                       metadata: DocumentMetadata) -> ParsedDocument:
        """Parse Markdown files with structure extraction."""
        # First parse as text
        parsed_doc = self._parse_text(file_path, strategy, metadata)
        
        if HAS_MARKDOWN_SUPPORT and strategy in [ParsingStrategy.BALANCED, ParsingStrategy.COMPREHENSIVE]:
            try:
                import markdown
                from markdown.extensions import toc
                
                # Convert to HTML for structure analysis
                md = markdown.Markdown(extensions=['toc', 'tables', 'fenced_code'])
                html_content = md.convert(parsed_doc.content)
                
                # Extract headers and structure
                sections = []
                lines = parsed_doc.content.split('\n')
                
                for line_num, line in enumerate(lines, 1):
                    if line.startswith('#'):
                        level = len(line) - len(line.lstrip('#'))
                        title = line.lstrip('# ').strip()
                        sections.append({
                            "type": "header",
                            "level": level,
                            "title": title,
                            "line_number": line_num
                        })
                
                parsed_doc.sections = sections
                
            except Exception as e:
                metadata.warnings.append(f"Markdown parsing failed: {e}")
        
        return parsed_doc
    
    def _parse_tex(self, file_path: Path, strategy: ParsingStrategy, 
                  metadata: DocumentMetadata) -> ParsedDocument:
        """Parse LaTeX/TeX files with advanced AST-based processing."""
        # Parse as text first for fallback
        parsed_doc = self._parse_text(file_path, strategy, metadata)
        
        if strategy in [ParsingStrategy.BALANCED, ParsingStrategy.COMPREHENSIVE]:
            try:
                # Use advanced AST-based LaTeX parsing
                content_chunks, preamble_metadata = self._parse_latex_ast(parsed_doc.content, file_path.parent)
                
                # Convert chunks to sections
                sections = []
                tables = []
                
                for chunk in content_chunks:
                    section_data = {
                        "type": chunk.get('type', 'paragraph'),
                        "content": chunk.get('content', ''),
                        "parent_section": chunk.get('parent_section', ''),
                        "hierarchy_path": chunk.get('metadata', {}).get('hierarchy_path', []),
                        "labels": chunk.get('metadata', {}).get('labels', []),
                        "citations": chunk.get('metadata', {}).get('citations', []),
                        "refs": chunk.get('metadata', {}).get('refs', [])
                    }
                    
                    # Separate tables from regular sections
                    if chunk.get('type') in ['table', 'tabular', 'longtable']:
                        tables.append(section_data)
                    else:
                        sections.append(section_data)
                
                parsed_doc.sections = sections
                parsed_doc.tables = tables
                
                # Update metadata with preamble information
                if preamble_metadata:
                    metadata.title = preamble_metadata.get('title')
                    metadata.author = preamble_metadata.get('author')
                    metadata.custom_properties['latex_preamble'] = preamble_metadata
                
            except Exception as e:
                # Fallback to basic parsing
                metadata.warnings.append(f"Advanced LaTeX parsing failed: {e}")
                self._parse_tex_basic(parsed_doc)
        
        return parsed_doc
    
    def _parse_html(self, file_path: Path, strategy: ParsingStrategy, 
                   metadata: DocumentMetadata) -> ParsedDocument:
        """Parse HTML files with tag extraction."""
        # Parse as text first
        parsed_doc = self._parse_text(file_path, strategy, metadata)
        
        try:
            from bs4 import BeautifulSoup
            
            soup = BeautifulSoup(parsed_doc.content, 'html.parser')
            
            # Extract text content
            text_content = soup.get_text(separator='\n', strip=True)
            parsed_doc.content = text_content
            
            # Extract structure if comprehensive
            if strategy == ParsingStrategy.COMPREHENSIVE:
                sections = []
                
                # Extract headers
                for tag in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6']):
                    sections.append({
                        "type": "html_header",
                        "tag": tag.name,
                        "text": tag.get_text(strip=True),
                        "level": int(tag.name[1])
                    })
                
                # Extract links
                links = []
                for link in soup.find_all('a', href=True):
                    links.append({
                        "text": link.get_text(strip=True),
                        "url": link['href'],
                        "type": "link"
                    })
                
                parsed_doc.sections = sections
                parsed_doc.links = links
                
        except ImportError:
            metadata.warnings.append("BeautifulSoup not available for HTML parsing")
        except Exception as e:
            metadata.warnings.append(f"HTML parsing failed: {e}")
        
        return parsed_doc
    
    def _parse_rtf(self, file_path: Path, strategy: ParsingStrategy, 
                  metadata: DocumentMetadata) -> ParsedDocument:
        """Parse RTF files (basic text extraction)."""
        # For now, treat as text - could be enhanced with RTF-specific library
        parsed_doc = self._parse_text(file_path, strategy, metadata)
        metadata.warnings.append("RTF parsing is basic - formatting may be lost")
        return parsed_doc
    
    def _analyze_structure(self, parsed_doc: ParsedDocument) -> None:
        """Analyze document structure and add metadata."""
        content = parsed_doc.content
        
        # Basic statistics
        parsed_doc.metadata.character_count = len(content)
        parsed_doc.metadata.word_count = len(content.split())
        
        # Language detection (simple heuristic)
        if len(content) > 100:
            # Simple language detection based on common words
            english_indicators = ['the', 'and', 'or', 'but', 'in', 'on', 'at', 'to', 'for']
            words = content.lower().split()
            english_score = sum(1 for word in words[:100] if word in english_indicators)
            
            if english_score > 5:
                parsed_doc.metadata.language = "en"
            else:
                parsed_doc.metadata.language = "unknown"
    
    def _parse_latex_ast(self, content: str, source_directory: Path) -> tuple:
        """Parse LaTeX content using AST-based approach from modules/chunker.py."""
        try:
            # Import LaTeX parsing dependencies
            from pylatexenc.latexwalker import LatexWalker, LatexCharsNode, LatexMacroNode, LatexEnvironmentNode, LatexCommentNode
            from pylatexenc.macrospec import LatexContextDb, SpecialsSpec
            import re
            import os
            
            # Resolve input files
            resolved_content = self._resolve_latex_inputs(content, source_directory)
            
            # Create LaTeX context
            db = LatexContextDb()
            db.add_context_category(
                'comments',
                specials=[SpecialsSpec('%')]
            )
            
            # Parse LaTeX AST
            lw = LatexWalker(resolved_content, latex_context=db)
            nodelist, _, _ = lw.get_latex_nodes()
            
            # Separate preamble from document body
            preamble_nodes = []
            body_nodelist = []
            doc_env_node = None
            
            for node in nodelist:
                if node.isNodeType(LatexEnvironmentNode) and node.environmentname == 'document':
                    doc_env_node = node
                    break
                else:
                    preamble_nodes.append(node)
            
            # Extract preamble metadata
            preamble_metadata = self._extract_latex_preamble_metadata(preamble_nodes)
            
            # Get document body
            if doc_env_node:
                body_nodelist = doc_env_node.nodelist
            else:
                start_index = len(preamble_nodes)
                body_nodelist = nodelist[start_index:]
            
            # Parse body content into chunks
            content_chunks = self._parse_latex_body(body_nodelist, [])
            
            return content_chunks, preamble_metadata
            
        except ImportError:
            raise DocumentParsingError("LaTeX AST parsing requires pylatexenc. Install with: pip install pylatexenc")
        except Exception as e:
            raise DocumentParsingError(f"LaTeX AST parsing failed: {e}")
    
    def _resolve_latex_inputs(self, content: str, base_dir: Path) -> str:
        """Resolve \input and \include commands in LaTeX content."""
        import re
        
        input_pattern = re.compile(r'\\(?:input|include)\s*\{([^}]+)\}')
        
        def replacer(match):
            filename = match.group(1).strip()
            if not filename.endswith('.tex'):
                filename += '.tex'
            filepath = base_dir / filename
            
            if filepath.exists():
                try:
                    with open(filepath, 'r', encoding='utf-8') as f:
                        return self._resolve_latex_inputs(f.read(), filepath.parent)
                except Exception as e:
                    logger.warning(f"Failed to read input file {filepath}: {e}")
                    return match.group(0)
            else:
                logger.warning(f"Input file not found: {filepath}")
                return match.group(0)
        
        return input_pattern.sub(replacer, content)
    
    def _extract_latex_preamble_metadata(self, preamble_nodes) -> dict:
        """Extract metadata from LaTeX preamble."""
        from pylatexenc.latexwalker import LatexMacroNode
        
        metadata = {
            'title': 'Untitled Document',
            'author': 'Unknown Author',
            'date': r'\today'
        }
        
        for node in preamble_nodes:
            if node.isNodeType(LatexMacroNode) and node.macroname in ['title', 'author', 'date']:
                if node.nodeargs:
                    metadata[node.macroname] = self._get_latex_node_text(node.nodeargs[0])
        
        return metadata
    
    def _parse_latex_body(self, nodelist, current_hierarchy) -> list:
        """Parse LaTeX document body into structured chunks."""
        from pylatexenc.latexwalker import LatexEnvironmentNode, LatexMacroNode, LatexCommentNode
        
        chunks = []
        buffer = []
        
        for node in nodelist:
            # Handle environments
            if node.isNodeType(LatexEnvironmentNode):
                # Process any buffered content first
                chunks.extend(self._process_latex_buffer(buffer, current_hierarchy))
                buffer = []
                chunks.append(self._create_latex_environment_chunk(node, current_hierarchy))
                continue
            
            # Handle section commands
            is_section_command = (
                node.isNodeType(LatexMacroNode) and
                node.macroname in ['section', 'subsection', 'subsubsection', 'paragraph']
            )
            
            if is_section_command:
                # Process buffered content
                chunks.extend(self._process_latex_buffer(buffer, current_hierarchy))
                buffer = []
                
                # Update hierarchy
                level = {'section': 1, 'subsection': 2, 'subsubsection': 3, 'paragraph': 4}[node.macroname]
                section_title = self._get_latex_node_text(node.nodeargs[0]).strip() if node.nodeargs else f"Untitled {node.macroname}"
                current_hierarchy = current_hierarchy[:level - 1] + [section_title]
                continue
            
            # Skip comments, add everything else to buffer
            if not node.isNodeType(LatexCommentNode):
                buffer.append(node)
        
        # Process any remaining buffer content
        chunks.extend(self._process_latex_buffer(buffer, current_hierarchy))
        return chunks
    
    def _process_latex_buffer(self, buffer, hierarchy) -> list:
        """Process buffered LaTeX nodes into chunks."""
        if not buffer:
            return []
        
        content = ''.join([n.latex_verbatim() for n in buffer]).strip()
        if not content:
            return []
        
        parent_section_str = ' -> '.join(hierarchy) if hierarchy else 'Preamble'
        
        # Extract metadata from buffer
        metadata = self._extract_latex_metadata(buffer)
        metadata['hierarchy_path'] = hierarchy.copy() if hierarchy else ['Preamble']
        
        chunk_data = {
            'type': 'paragraph',
            'content': content,
            'parent_section': parent_section_str,
            'metadata': metadata
        }
        
        return [chunk_data]
    
    def _create_latex_environment_chunk(self, node, hierarchy) -> dict:
        """Create a chunk for a LaTeX environment."""
        parent_section_str = ' -> '.join(hierarchy) if hierarchy else 'Preamble'
        
        # Extract caption if present
        caption = self._find_latex_caption(node.nodelist or [])
        
        metadata = {
            'hierarchy_path': hierarchy.copy(),
            'caption': caption,
            **self._extract_latex_metadata(node.nodelist or [])
        }
        
        return {
            'type': node.environmentname,
            'content': node.latex_verbatim(),
            'parent_section': parent_section_str,
            'metadata': metadata
        }
    
    def _extract_latex_metadata(self, nodelist) -> dict:
        """Extract LaTeX metadata (labels, citations, refs) from node list."""
        from pylatexenc.latexwalker import LatexMacroNode, LatexEnvironmentNode
        
        labels, citations, refs = [], [], []
        
        for node in nodelist:
            if node.isNodeType(LatexMacroNode):
                if node.macroname == 'label' and node.nodeargs:
                    labels.append(self._get_latex_node_text(node.nodeargs[0]))
                elif node.macroname in ('cite', 'citep', 'citet') and node.nodeargs:
                    citations.extend(self._get_latex_node_text(node.nodeargs[0]).split(','))
                elif node.macroname == 'ref' and node.nodeargs:
                    refs.append(self._get_latex_node_text(node.nodeargs[0]))
            elif node.isNodeType(LatexEnvironmentNode) and node.nodelist:
                child_metadata = self._extract_latex_metadata(node.nodelist)
                labels.extend(child_metadata['labels'])
                citations.extend(child_metadata['citations'])
                refs.extend(child_metadata['refs'])
        
        return {
            'labels': [l.strip() for l in labels],
            'citations': [c.strip() for c in citations],
            'refs': [r.strip() for r in refs]
        }
    
    def _get_latex_node_text(self, node) -> str:
        """Extract text content from a LaTeX node."""
        return node.latex_verbatim().strip('{}')
    
    def _find_latex_caption(self, nodelist) -> str:
        """Find caption in LaTeX node list."""
        from pylatexenc.latexwalker import LatexMacroNode
        
        for node in nodelist:
            if node.isNodeType(LatexMacroNode) and node.macroname == 'caption' and node.nodeargs:
                return self._get_latex_node_text(node.nodeargs[0])
        return None
    
    def _parse_tex_basic(self, parsed_doc: ParsedDocument) -> None:
        """Basic LaTeX parsing fallback method."""
        sections = []
        lines = parsed_doc.content.split('\n')
        
        for line_num, line in enumerate(lines, 1):
            stripped = line.strip()
            
            # Detect sections
            if stripped.startswith('\\section{') or stripped.startswith('\\subsection{') or stripped.startswith('\\subsubsection{'):
                command = stripped.split('{')[0][1:]  # Remove \\
                title = stripped.split('{', 1)[1].rsplit('}', 1)[0] if '{' in stripped else ""
                
                sections.append({
                    "type": "latex_section",
                    "command": command,
                    "title": title,
                    "line_number": line_num
                })
        
        parsed_doc.sections = sections
    
    def _calculate_quality_score(self, parsed_doc: ParsedDocument) -> float:
        """Calculate a quality score for the parsed document."""
        scores = []
        
        # Content length score
        content_length = len(parsed_doc.content.strip())
        if content_length > 1000:
            scores.append(1.0)
        elif content_length > 100:
            scores.append(0.7)
        elif content_length > 10:
            scores.append(0.4)
        else:
            scores.append(0.1)
        
        # Structure score
        if parsed_doc.sections:
            scores.append(0.8)
        else:
            scores.append(0.5)
        
        # Warning penalty
        warning_penalty = min(0.2 * len(parsed_doc.metadata.warnings), 0.5)
        
        # Calculate final score
        base_score = sum(scores) / len(scores) if scores else 0.5
        final_score = max(0.0, base_score - warning_penalty)
        
        return final_score


class DocumentParsingError(Exception):
    """Custom exception for document parsing errors."""
    pass


# Factory functions for easy usage
def create_parser(strategy: ParsingStrategy = ParsingStrategy.BALANCED, 
                 **kwargs) -> DocumentParser:
    """Create a document parser with specified strategy."""
    return DocumentParser(default_strategy=strategy, **kwargs)


async def parse_document_async(file_path: Union[str, Path], 
                             strategy: ParsingStrategy = ParsingStrategy.BALANCED) -> ParsedDocument:
    """Convenience function for async document parsing."""
    parser = DocumentParser(default_strategy=strategy)
    return await parser.parse_document_async(file_path, strategy)